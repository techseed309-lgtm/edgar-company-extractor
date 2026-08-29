"""Build the client-facing Excel workbook and CSV exports."""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LINK_FONT = Font(color="0563C1", underline="single")

COMPANY_HEADERS = {
    "cik": "CIK", "name": "Company Name", "ticker": "Ticker", "exchange": "Exchange",
    "sic_code": "SIC Code", "industry": "Industry", "phone": "Phone", "street": "Street",
    "city": "City", "state": "State", "zip_code": "ZIP",
    "state_of_incorporation": "State of Inc.", "fiscal_year_end": "FY End",
    "filings_matched": "Filings Found", "missing_fields": "Missing Fields",
    "edgar_url": "EDGAR Profile",
}
FILING_HEADERS = {
    "cik": "CIK", "company": "Company Name", "ticker": "Ticker", "form": "Form Type",
    "filing_date": "Filing Date", "report_date": "Report Date",
    "accession_number": "Accession Number", "document_url": "Document URL",
    "size_kb": "Size (KB)",
}


def _style_sheet(worksheet, link_columns: set[str] = frozenset()) -> None:
    """Apply header styling, freeze panes, autofilter and sane column widths."""
    if worksheet.max_row < 1:
        return

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    worksheet.freeze_panes = "A2"
    if worksheet.max_row > 1:
        worksheet.auto_filter.ref = worksheet.dimensions

    for column_cells in worksheet.columns:
        letter = get_column_letter(column_cells[0].column)
        longest = max((len(str(c.value)) for c in column_cells if c.value is not None), default=8)
        worksheet.column_dimensions[letter].width = min(max(longest + 2, 10), 46)

        if column_cells[0].value in link_columns:
            for cell in column_cells[1:]:
                if cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.font = LINK_FONT


def _summary_frame(companies: pd.DataFrame, filings: pd.DataFrame, stats: dict) -> pd.DataFrame:
    """Flat key/value sheet describing the run, plus simple breakdowns."""
    rows = [
        ("Run completed", stats["finished_at"]),
        ("Duration (seconds)", stats["duration_seconds"]),
        ("Identifiers requested", stats["requested"]),
        ("Companies extracted", len(companies)),
        ("Companies failed", stats["failed"]),
        ("Success rate", f"{stats['success_rate']:.1%}"),
        ("Filings extracted", len(filings)),
        ("HTTP downloads", stats["downloads"]),
        ("Served from cache", stats["cache_hits"]),
        ("Retries performed", stats["retries"]),
        ("", ""),
        ("--- Companies by state ---", ""),
    ]
    if not companies.empty:
        for state, count in companies["State"].replace("", "Unknown").value_counts().items():
            rows.append((str(state), int(count)))
        rows.append(("", ""))
        rows.append(("--- Filings by form type ---", ""))
        if not filings.empty:
            for form, count in filings["Form Type"].value_counts().items():
                rows.append((str(form), int(count)))
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def write_workbook(companies, filings, errors, stats, output_path: Path) -> Path:
    """Write the four-sheet deliverable. Returns the path written."""
    output_path.parent.mkdir(parents=True, exist_ok=True)

    company_df = pd.DataFrame([c.to_row() for c in companies]).rename(columns=COMPANY_HEADERS)
    if not company_df.empty:
        company_df = company_df[[v for v in COMPANY_HEADERS.values() if v in company_df.columns]]

    filing_df = pd.DataFrame([f.to_row() for f in filings]).rename(columns=FILING_HEADERS)
    if not filing_df.empty:
        filing_df = filing_df[[v for v in FILING_HEADERS.values() if v in filing_df.columns]]

    error_df = pd.DataFrame(errors, columns=["Input Identifier", "Reason", "Detail"])
    if error_df.empty:
        error_df = pd.DataFrame([["-", "No errors", "All identifiers resolved successfully"]],
                                columns=["Input Identifier", "Reason", "Detail"])

    summary_df = _summary_frame(company_df, filing_df, stats)

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        company_df.to_excel(writer, sheet_name="Companies", index=False)
        filing_df.to_excel(writer, sheet_name="Filings", index=False)
        error_df.to_excel(writer, sheet_name="Run Log", index=False)

        _style_sheet(writer.sheets["Summary"])
        _style_sheet(writer.sheets["Companies"], link_columns={"EDGAR Profile"})
        _style_sheet(writer.sheets["Filings"], link_columns={"Document URL"})
        _style_sheet(writer.sheets["Run Log"])

    log.info("Wrote workbook: %s", output_path)
    return output_path


def write_csvs(companies, filings, output_dir: Path, stem: str) -> list[Path]:
    """Also emit plain CSVs; some clients ask for them explicitly."""
    output_dir.mkdir(parents=True, exist_ok=True)
    written = []
    pairs = [
        (pd.DataFrame([c.to_row() for c in companies]).rename(columns=COMPANY_HEADERS), "companies"),
        (pd.DataFrame([f.to_row() for f in filings]).rename(columns=FILING_HEADERS), "filings"),
    ]
    for frame, suffix in pairs:
        path = output_dir / f"{stem}_{suffix}.csv"
        frame.to_csv(path, index=False, encoding="utf-8-sig")
        written.append(path)
        log.info("Wrote CSV: %s", path)
    return written
